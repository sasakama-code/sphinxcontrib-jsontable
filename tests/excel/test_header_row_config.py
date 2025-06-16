"""Phase 2: Header Row Configuration機能のTDDテスト.

Task 2.3: `:header-row:` オプション実装のテスト
- 行番号指定機能
- 複数行ヘッダー対応準備
- 指定行からのヘッダー取得
- ヘッダー名の正規化
"""

import os
import shutil
import tempfile

import pytest
from sphinx.util.docutils import docutils_namespace

# Excel対応がある場合のみテストを実行
try:
    from sphinxcontrib.jsontable.directives import JsonTableDirective
    from sphinxcontrib.jsontable.excel_data_loader import ExcelDataLoader

    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


def create_mock_state_machine(srcdir="/tmp"):
    """Create a mock state machine for testing JsonTableDirective."""

    class MockReporter:
        def warning(self, msg, *args, **kwargs):
            pass

        def error(self, msg, *args, **kwargs):
            pass

        def info(self, msg, *args, **kwargs):
            pass

    class MockConfig:
        def __init__(self):
            self.jsontable_max_rows = 1000

    class MockEnv:
        def __init__(self, srcdir):
            self.config = MockConfig()
            self.srcdir = srcdir

    class MockSettings:
        def __init__(self, srcdir):
            self.env = MockEnv(srcdir)

    class MockDocument:
        def __init__(self, srcdir):
            self.settings = MockSettings(srcdir)

    class MockState:
        def __init__(self, srcdir):
            self.document = MockDocument(srcdir)

    class MockStateMachine:
        def __init__(self):
            self.reporter = MockReporter()

    return MockStateMachine(), MockState(srcdir)


class TestHeaderRowConfiguration:
    """Phase 2: Header Row Configuration機能のテスト."""

    def setup_method(self):
        """各テストメソッドの前に実行される。"""
        self.temp_dir = tempfile.mkdtemp()
        self.loader = ExcelDataLoader(self.temp_dir)

    def teardown_method(self):
        """各テストメソッドの後に実行される。"""
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def create_header_test_excel(self) -> str:
        """ヘッダー行設定テスト用のExcelファイルを作成。

        Returns:
            str: 作成されたファイルのパス
        """
        file_path = os.path.join(self.temp_dir, "header_test.xlsx")

        # 複数行でヘッダーが異なる位置にあるデータを作成
        data = [
            ["メタデータ", "作成日: 2025-06-13", "", ""],  # Row 0: メタデータ行
            ["説明", "売上データの月次集計", "", ""],  # Row 1: 説明行
            ["", "", "", ""],  # Row 2: 空行
            ["商品名", "1月売上", "2月売上", "3月売上"],  # Row 3: ヘッダー行
            ["商品A", "100000", "120000", "110000"],  # Row 4: データ行
            ["商品B", "150000", "180000", "160000"],  # Row 5: データ行
            ["商品C", "80000", "90000", "85000"],  # Row 6: データ行
        ]

        # DataFrameではなく、直接openpyxlを使用してデータを書き込む
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # データを行ごとに書き込み
        for row_idx, row_data in enumerate(data, 1):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        wb.save(file_path)
        return file_path

    def test_header_row_specification(self):
        """指定行をヘッダーとして使用するテスト(未実装機能なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # 4行目(0ベースで3)をヘッダーとして指定
        result = self.loader.load_from_excel_with_header_row(excel_path, header_row=3)

        # 期待されるヘッダー
        expected_headers = ["商品名", "1月売上", "2月売上", "3月売上"]
        assert result["headers"] == expected_headers
        assert result["has_header"]

        # データ部分(header_row=3の場合、前の行は削除され、4,5,6行目がデータになる)
        expected_data = [
            ["商品A", "100000", "120000", "110000"],
            ["商品B", "150000", "180000", "160000"],
            ["商品C", "80000", "90000", "85000"],
        ]
        assert result["data"] == expected_data
        assert result["header_row"] == 3

    def test_header_row_with_range_combination(self):
        """ヘッダー行指定と範囲指定の組み合わせテスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # 4行目をヘッダー、A4:C7の範囲を指定
        result = self.loader.load_from_excel_with_header_row_and_range(
            excel_path, header_row=3, range_spec="A4:C7"
        )

        expected_headers = ["商品名", "1月売上", "2月売上"]
        assert result["headers"] == expected_headers

        # データ部分(5-7行目、A-C列)
        expected_data = [
            ["商品A", "100000", "120000"],
            ["商品B", "150000", "180000"],
            ["商品C", "80000", "90000"],
        ]
        assert result["data"] == expected_data

    def test_invalid_header_row_error(self):
        """無効なヘッダー行指定時のエラーテスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # 負のヘッダー行番号
        with pytest.raises(ValueError, match="Header row must be non-negative"):
            self.loader.load_from_excel_with_header_row(excel_path, header_row=-1)

        # 範囲外のヘッダー行番号
        with pytest.raises(ValueError, match="Header row 100 is out of range"):
            self.loader.load_from_excel_with_header_row(excel_path, header_row=100)

    def test_empty_header_row_handling(self):
        """空のヘッダー行の処理テスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # 3行目(空行)をヘッダーとして指定
        result = self.loader.load_from_excel_with_header_row(excel_path, header_row=2)

        # 空のヘッダーは自動生成されるべき(正規化後)
        expected_headers = ["column_1", "column_2", "column_3", "column_4"]
        assert result["headers"] == expected_headers
        assert result["has_header"]

    def test_directive_header_row_option(self):
        """JsonTableDirectiveの:header-row:オプションテスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # モックSphinx環境
        class MockConfig:
            jsontable_max_rows = 10000

        class MockEnv:
            def __init__(self, srcdir):
                self.srcdir = srcdir
                self.config = MockConfig()

        MockEnv(self.temp_dir)

        with docutils_namespace():
            # ヘッダー行指定付きディレクティブ
            mock_state_machine, mock_state = create_mock_state_machine(self.temp_dir)

            directive = JsonTableDirective(
                name="jsontable",
                arguments=[os.path.basename(excel_path)],
                options={"header": True, "header-row": 3},
                content=[],
                lineno=1,
                content_offset=0,
                block_text="",
                state=mock_state,
                state_machine=mock_state_machine,
            )
            directive.excel_loader = ExcelDataLoader(self.temp_dir)

            json_data = directive._load_json_data()

            # 4行目がヘッダーとして使用されることを確認
            assert len(json_data) == 3  # データ行数
            assert json_data[0]["商品名"] == "商品A"
            assert json_data[0]["1月売上"] == "100000"
            assert json_data[2]["商品名"] == "商品C"

    def test_directive_header_row_with_sheet_option(self):
        """ヘッダー行とシート指定の組み合わせテスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        class MockConfig:
            jsontable_max_rows = 10000

        class MockEnv:
            def __init__(self, srcdir):
                self.srcdir = srcdir
                self.config = MockConfig()

        MockEnv(self.temp_dir)

        with docutils_namespace():
            # シートとヘッダー行の両方を指定
            mock_state_machine, mock_state = create_mock_state_machine(self.temp_dir)

            directive = JsonTableDirective(
                name="jsontable",
                arguments=[os.path.basename(excel_path)],
                options={"header": True, "sheet": "Sheet1", "header-row": 3},
                content=[],
                lineno=1,
                content_offset=0,
                block_text="",
                state=mock_state,
                state_machine=mock_state_machine,
            )
            directive.excel_loader = ExcelDataLoader(self.temp_dir)

            json_data = directive._load_json_data()

            # Sheet1の4行目がヘッダーとして使用されることを確認
            assert len(json_data) == 3
            assert json_data[0]["商品名"] == "商品A"

    def test_header_row_normalization(self):
        """ヘッダー名の正規化テスト(未実装なので失敗する)。"""
        excel_path = self.create_header_test_excel()

        # ヘッダー名に空白や特殊文字が含まれる場合の正規化
        result = self.loader.load_from_excel_with_header_row(excel_path, header_row=3)

        # ヘッダー名が適切に正規化されることを確認
        normalized_headers = self.loader._normalize_header_names(result["headers"])

        # 期待される正規化(例:空白トリム、重複回避)
        expected_normalized = ["商品名", "1月売上", "2月売上", "3月売上"]
        assert normalized_headers == expected_normalized

    def test_header_row_validation(self):
        """ヘッダー行の検証ロジックテスト(未実装なので失敗する)。"""
        # ヘッダー行番号が数値でない場合
        with pytest.raises(TypeError, match="Header row must be an integer"):
            self.loader._validate_header_row("invalid")

        # ヘッダー行番号がNoneの場合
        assert self.loader._validate_header_row(None) is None  # 自動検出モード


if __name__ == "__main__":
    # スタンドアロンテスト実行
    pytest.main([__file__, "-v"])
